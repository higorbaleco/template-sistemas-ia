from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path("/Users/higorplens/Documents/Codex/2026-10-05/cara-seguinte-preciso-organizar-melhor-nosso/outputs/sistema-comercial-onboarding-roi.xlsx")

NAVY = "172554"
TEAL = "0F766E"
GREEN = "16A34A"
AMBER = "F59E0B"
RED = "DC2626"
LIGHT = "F8FAFC"
MID = "E2E8F0"
PALE_BLUE = "DBEAFE"
PALE_GREEN = "DCFCE7"
PALE_YELLOW = "FEF3C7"
WHITE = "FFFFFF"
BLUE_INPUT = "0000FF"
BLACK = "111827"

thin = Side(style="thin", color=MID)


def title(ws, text, subtitle=None):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    if subtitle:
        ws.merge_cells("A2:H2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(name="Arial", size=10, color="475569")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 32


def style_headers(ws, row, start, end):
    for cell in ws.iter_cols(min_col=start, max_col=end, min_row=row, max_row=row):
        c = cell[0]
        c.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    ws.row_dimensions[row].height = 42


def style_body(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            c.font = Font(name="Arial", size=10, color=BLACK)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=thin)


def input_cells(ws, cells):
    for ref in cells:
        ws[ref].font = Font(name="Arial", size=10, color=BLUE_INPUT)
        ws[ref].fill = PatternFill("solid", fgColor=PALE_YELLOW)


def add_table(ws, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def widths(ws, mapping):
    for col, width in mapping.items():
        ws.column_dimensions[col].width = width


def dropdown(ws, cell_range, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)


wb = Workbook()
ws = wb.active
ws.title = "LEIA-ME"
title(ws, "Sistema comercial, onboarding e ROI", "Modelo para Excel ou importação no Google Sheets. Edite células amarelas com texto azul; fórmulas ficam em preto.")

guide = [
    ("1", "Pré-venda", "Use Roteiro Pré-venda durante a ligação e registre uma linha em Oportunidades."),
    ("2", "Decisão de fit", "A planilha calcula nota, classificação e exposição financeira preliminar."),
    ("3", "Escopo e preço", "Preencha Escopo Setup para estimar horas, custo, preço mínimo, preço sugerido e prazo."),
    ("4", "Fechamento", "Ao ganhar, aprofunde o briefing na aba Onboarding; não pergunte novamente o que já foi coletado."),
    ("5", "Plano do cliente", "Registre um agente por linha em Plano do Cliente, com função, limites, integrações e métricas."),
    ("6", "Operação mensal", "Estime horas, tokens, custo e mensalidade em Operação Mensal."),
    ("7", "Prova de valor", "Preencha Relatório ROI mensalmente com linha de base e resultados observados."),
]
ws.append([])
ws.append(["Etapa", "Área", "Como usar"])
for item in guide:
    ws.append(item)
style_headers(ws, 4, 1, 3)
style_body(ws, 5, 11, 1, 3)
widths(ws, {"A": 10, "B": 24, "C": 95})
ws.freeze_panes = "A4"

ws = wb.create_sheet("Parâmetros")
title(ws, "Parâmetros comerciais e operacionais", "Valores de exemplo. Substitua pelos custos, preços e capacidade reais da Avraham.")
params = [
    ("Custo interno carregado por hora (R$)", 120, "Inclua salário/pró-labore, encargos, ferramentas e ociosidade."),
    ("Valor comercial por hora (R$)", 300, "Referência para preço sugerido de setup."),
    ("Margem bruta alvo", 0.60, "Usada para calcular preço mínimo saudável."),
    ("Gestão do projeto sobre horas", 0.15, "Percentual adicional para coordenação."),
    ("Contingência sobre horas", 0.10, "Reserva para risco e variação de escopo."),
    ("Capacidade semanal por projeto (h)", 20, "Capacidade efetiva, não jornada total."),
    ("Captura de valor mensal", 0.10, "Parcela do impacto conservador usada como referência de preço."),
    ("Margem sobre tokens/infra", 0.30, "Markup operacional sobre custos variáveis."),
    ("Câmbio USD/BRL", 5.50, "Exemplo; atualizar conforme política financeira."),
    ("Modelo - entrada (USD/1M tokens)", 1.00, "Exemplo; substituir pelo modelo contratado."),
    ("Modelo - saída (USD/1M tokens)", 4.00, "Exemplo; substituir pelo modelo contratado."),
    ("Descoberta e diagnóstico (h)", 4, "Base por cliente."),
    ("Mapeamento do processo (h)", 6, "Base por processo principal."),
    ("Criação por agente (h)", 10, "Arquitetura e instruções centrais."),
    ("Fluxo/intenção (h)", 1.5, "Estimativa por fluxo relevante."),
    ("Fonte de conhecimento (h)", 0.5, "Tratamento por documento ou lote simples."),
    ("Integração simples (h)", 8, "API conhecida, baixo risco e poucos campos."),
    ("Integração complexa (h)", 24, "Regras, autenticação ou escrita complexa."),
    ("Canal (h)", 4, "Configuração por canal."),
    ("Ação customizada (h)", 6, "Ferramenta ou função específica."),
    ("Cenário de teste (h)", 0.5, "Criação, execução e registro."),
    ("Treinamento/handoff (h)", 3, "Por sessão."),
]
ws.append([])
ws.append(["Parâmetro", "Valor", "Observação"])
for row in params:
    ws.append(row)
style_headers(ws, 4, 1, 3)
style_body(ws, 5, 26, 1, 3)
for row in range(5, 27):
    input_cells(ws, [f"B{row}"])
    ws[f"B{row}"].comment = Comment(ws[f"C{row}"].value, "Avraham")
for row in [7, 8, 9, 11, 12]:
    ws[f"B{row}"].number_format = "0.0%"
for row in [5, 6, 13, 14, 15]:
    ws[f"B{row}"].number_format = 'R$ #,##0.00;[Red](R$ #,##0.00);-'
widths(ws, {"A": 42, "B": 20, "C": 72})
ws.freeze_panes = "A5"

ws = wb.create_sheet("Roteiro Pré-venda")
title(ws, "Roteiro da ligação de descoberta", "Perguntas para descobrir fit, impacto, escopo provável e argumentos da proposta.")
questions = [
    ("Contexto", "O que a empresa vende, para quem e por qual canal?", "Contexto comercial", "Texto"),
    ("Contexto", "Qual processo vocês gostariam de melhorar com IA?", "Caso de uso", "Texto"),
    ("Prioridade", "Por que isso virou prioridade agora?", "Urgência", "Texto"),
    ("Prioridade", "O que acontece se nada mudar nos próximos seis meses?", "Custo de inação", "Texto"),
    ("Dor", "Onde o processo mais trava, atrasa ou perde qualidade?", "Dor principal", "Texto"),
    ("Dor", "Conte um caso recente em que isso custou tempo, venda ou satisfação.", "Evidência", "Texto"),
    ("Volume", "Quantos contatos ou solicitações entram por mês?", "Escala", "Número"),
    ("Impacto", "Qual é o ticket médio e a conversão atual?", "Receita exposta", "Moeda + %"),
    ("Impacto", "Quantos contatos ficam sem resposta ou sem follow-up?", "Abandono", "% ou número"),
    ("Impacto", "Quanto tempo leva a primeira resposta?", "Velocidade", "Minutos"),
    ("Impacto", "Quantas pessoas e horas mensais são gastas nesse processo?", "Custo de trabalho", "Número"),
    ("Processo", "Descreva o caminho do primeiro contato até a conclusão.", "Complexidade", "Texto"),
    ("Processo", "Quais decisões são regras e quais dependem de julgamento humano?", "Automatização possível", "Texto"),
    ("Tecnologia", "Quais sistemas, canais e integrações fazem parte do processo?", "Esforço técnico", "Lista"),
    ("Tecnologia", "Existe API, documentação e apoio técnico do cliente?", "Prontidão técnica", "Sim/Não + detalhe"),
    ("Conhecimento", "Quais materiais e conversas reais existem?", "Prontidão de conteúdo", "Lista"),
    ("Risco", "O que a IA jamais poderia fazer ou prometer?", "Limites", "Texto"),
    ("Risco", "Quando a conversa precisa ir para uma pessoa?", "Handoff", "Texto"),
    ("Compra", "Quem aprova e quem será o dono interno?", "Governança", "Texto"),
    ("Compra", "Existe data e faixa de investimento previstas?", "Fit comercial", "Data + faixa"),
    ("Métrica", "Que resultado provaria sucesso em 90 dias?", "Critério de sucesso", "Texto + número"),
    ("Métrica", "O cliente aceita medir uma linha de base antes do go-live?", "Mensurabilidade", "Sim/Não"),
]
ws.append([])
ws.append(["Bloco", "Pergunta", "O que revela", "Formato da resposta"])
for row in questions:
    ws.append(row)
style_headers(ws, 4, 1, 4)
style_body(ws, 5, 26, 1, 4)
add_table(ws, "A4:D26", "tblRoteiroPreVenda")
widths(ws, {"A": 18, "B": 70, "C": 28, "D": 22})
ws.freeze_panes = "A5"

ws = wb.create_sheet("Oportunidades")
title(ws, "Registro e qualificação de oportunidades", "Uma linha por oportunidade. Notas de 0 a 5; bloqueador crítico força a classificação Sem fit.")
headers = [
    "ID", "Data", "Cliente", "Vendedor", "Segmento", "Contatos/mês", "Ticket médio (R$)",
    "Conversão atual", "% sem follow-up", "1ª resposta (min)", "Dor e prioridade (0-5)",
    "Impacto econômico (0-5)", "Volume/repetição (0-5)", "Prontidão operacional (0-5)",
    "Fit comercial (0-5)", "Bloqueador crítico?", "Nota de fit", "Classificação", "Receita exposta/mês",
    "Dor principal", "Resultado desejado", "Processo atual", "Sistemas/canais", "Decisor/dono", "Faixa de investimento",
    "Urgência/data", "Próximo passo", "Status", "Tese para proposta", "Motivo sem fit/perda"
]
ws.append([])
ws.append(headers)
example = [
    "OP-001", "05/10/2026", "Cliente Exemplo", "Closer", "Serviços", 1200, 2500, 0.08, 0.28, 180,
    5, 4, 5, 3, 4, "Não", "=K5*6+L5*5+M5*3+N5*3+O5*3",
    '=IF(P5="Sim","Sem fit",IF(Q5>=75,"Fit forte",IF(Q5>=55,"Fit condicionado","Sem fit")))',
    "=F5*I5*H5*G5", "Leads sem continuidade", "Proteger e qualificar 24/7", "Atendimento manual sem cadência",
    "WhatsApp + CRM", "Diretor comercial", "R$ 8 mil a R$ 15 mil", "30 dias", "Agendar proposta", "Fit forte",
    "Proteger resposta, qualificação e follow-up antes de ampliar.", ""
]
ws.append(example)
for row in range(6, 106):
    ws.cell(row, 17, f'=IF(C{row}="","",K{row}*6+L{row}*5+M{row}*3+N{row}*3+O{row}*3)')
    ws.cell(row, 18, f'=IF(C{row}="","",IF(P{row}="Sim","Sem fit",IF(Q{row}>=75,"Fit forte",IF(Q{row}>=55,"Fit condicionado","Sem fit"))))')
    ws.cell(row, 19, f'=IF(C{row}="","",F{row}*I{row}*H{row}*G{row})')
style_headers(ws, 4, 1, len(headers))
style_body(ws, 5, 105, 1, len(headers))
input_cols = list(range(1, 17)) + list(range(20, 31))
for row in range(5, 106):
    for col in input_cols:
        c = ws.cell(row, col)
        c.font = Font(name="Arial", size=10, color=BLUE_INPUT)
        c.fill = PatternFill("solid", fgColor=PALE_YELLOW)
for row in range(5, 106):
    ws.cell(row, 8).number_format = "0.0%"
    ws.cell(row, 9).number_format = "0.0%"
    ws.cell(row, 7).number_format = 'R$ #,##0.00'
    ws.cell(row, 19).number_format = 'R$ #,##0.00'
dropdown(ws, "K5:O105", ["0", "1", "2", "3", "4", "5"])
dropdown(ws, "P5:P105", ["Não", "Sim"])
dropdown(ws, "AB5:AB105", ["Novo lead", "Descoberta agendada", "Descoberta realizada", "Fit forte", "Fit condicionado", "Sem fit", "Validação técnica", "Proposta agendada", "Proposta enviada", "Negociação", "Ganho", "Perdido", "Nutrição"])
ws.conditional_formatting.add("Q5:Q105", ColorScaleRule(start_type="num", start_value=0, start_color=RED, mid_type="num", mid_value=55, mid_color=AMBER, end_type="num", end_value=100, end_color=GREEN))
add_table(ws, "A4:AD105", "tblOportunidades")
ws.auto_filter.ref = "A4:AD105"
ws.freeze_panes = "F5"
for col in range(1, 31):
    ws.column_dimensions[ws.cell(4, col).column_letter].width = 17
for col in [20, 21, 22, 23, 24, 27, 29, 30]:
    ws.column_dimensions[ws.cell(4, col).column_letter].width = 30

ws = wb.create_sheet("Escopo Setup")
title(ws, "Estimativa de setup", "Quantifique o escopo. O preço mínimo protege margem; o sugerido usa a maior referência entre custo e hora comercial.")
headers = [
    "ID oportunidade", "Cliente", "Complexidade", "Agentes", "Fluxos", "Fontes conhecimento",
    "Integrações simples", "Integrações complexas", "Canais", "Ações customizadas", "Cenários teste",
    "Sessões treinamento", "Horas migração/outros", "Descoberta (h)", "Mapeamento (h)", "Agentes (h)",
    "Fluxos (h)", "Conhecimento (h)", "Integrações (h)", "Canais (h)", "Customização (h)",
    "Testes (h)", "Treinamento (h)", "Subtotal (h)", "Gestão (h)", "Contingência (h)",
    "Total (h)", "Custo interno", "Preço mínimo", "Preço sugerido", "Margem sugerida", "Prazo (semanas)"
]
ws.append([])
ws.append(headers)
inputs = ["OP-001", "Cliente Exemplo", 1.00, 1, 6, 8, 1, 0, 1, 1, 12, 1, 0]
ws.append(inputs)
for row in range(5, 55):
    ws.cell(row, 14, f'=IF(B{row}="","",Parâmetros!$B$16*C{row})')
    ws.cell(row, 15, f'=IF(B{row}="","",Parâmetros!$B$17*C{row})')
    ws.cell(row, 16, f'=IF(B{row}="","",D{row}*Parâmetros!$B$18*C{row})')
    ws.cell(row, 17, f'=IF(B{row}="","",E{row}*Parâmetros!$B$19*C{row})')
    ws.cell(row, 18, f'=IF(B{row}="","",F{row}*Parâmetros!$B$20*C{row})')
    ws.cell(row, 19, f'=IF(B{row}="","",(G{row}*Parâmetros!$B$21+H{row}*Parâmetros!$B$22)*C{row})')
    ws.cell(row, 20, f'=IF(B{row}="","",I{row}*Parâmetros!$B$23*C{row})')
    ws.cell(row, 21, f'=IF(B{row}="","",J{row}*Parâmetros!$B$24*C{row})')
    ws.cell(row, 22, f'=IF(B{row}="","",K{row}*Parâmetros!$B$25*C{row})')
    ws.cell(row, 23, f'=IF(B{row}="","",L{row}*Parâmetros!$B$26*C{row})')
    ws.cell(row, 24, f'=IF(B{row}="","",SUM(N{row}:W{row})+M{row})')
    ws.cell(row, 25, f'=IF(B{row}="","",X{row}*Parâmetros!$B$8)')
    ws.cell(row, 26, f'=IF(B{row}="","",X{row}*Parâmetros!$B$9)')
    ws.cell(row, 27, f'=IF(B{row}="","",SUM(X{row}:Z{row}))')
    ws.cell(row, 28, f'=IF(B{row}="","",AA{row}*Parâmetros!$B$5)')
    ws.cell(row, 29, f'=IF(B{row}="","",AB{row}/(1-Parâmetros!$B$7))')
    ws.cell(row, 30, f'=IF(B{row}="","",MAX(AA{row}*Parâmetros!$B$6,AC{row}))')
    ws.cell(row, 31, f'=IF(B{row}="","",(AD{row}-AB{row})/AD{row})')
    ws.cell(row, 32, f'=IF(B{row}="","",AA{row}/Parâmetros!$B$10)')
style_headers(ws, 4, 1, len(headers))
style_body(ws, 5, 54, 1, len(headers))
for row in range(5, 55):
    for col in range(1, 14):
        ws.cell(row, col).font = Font(name="Arial", size=10, color=BLUE_INPUT)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    for col in [28, 29, 30]:
        ws.cell(row, col).number_format = 'R$ #,##0.00;[Red](R$ #,##0.00);-'
    ws.cell(row, 31).number_format = "0.0%"
dropdown(ws, "C5:C54", ["0.75", "1", "1.5", "2"])
add_table(ws, "A4:AF54", "tblEscopoSetup")
ws.freeze_panes = "D5"
for col in range(1, 33):
    ws.column_dimensions[ws.cell(4, col).column_letter].width = 17

ws = wb.create_sheet("Onboarding")
title(ws, "Briefing de onboarding", "Use após o fechamento. A coluna Evidência esperada evita respostas vagas.")
onboarding = [
    ("Objetivo", "Quais dores foram contratadas e em que ordem serão atacadas?", "Lista priorizada", "Obrigatório"),
    ("Objetivo", "Qual resultado deve ser alcançado em 30, 60 e 90 dias?", "Meta e prazo", "Obrigatório"),
    ("Baseline", "Qual é a linha de base de volume, tempo, conversão, custo e qualidade?", "Relatório ou amostra", "Obrigatório"),
    ("Processo", "Qual é o fluxo completo, incluindo exceções e caminhos alternativos?", "Fluxograma validado", "Obrigatório"),
    ("Agentes", "Quantos papéis distintos exigem contexto, permissões ou métricas diferentes?", "Lista de agentes", "Obrigatório"),
    ("Fluxos", "Quais são as principais intenções e o resultado esperado de cada uma?", "Matriz intenção x ação", "Obrigatório"),
    ("Qualificação", "Quais critérios qualificam, desqualificam ou priorizam um contato?", "Regras objetivas", "Se aplicável"),
    ("Regras", "Quais regras nunca podem ser quebradas e quem autoriza exceções?", "Política validada", "Obrigatório"),
    ("Limites", "O que a IA não pode dizer, decidir, consultar ou alterar?", "Lista de proibições", "Obrigatório"),
    ("Handoff", "Quando, para quem e com qual resumo a IA transfere?", "Matriz de escalonamento", "Obrigatório"),
    ("Integrações", "O que cada sistema permite consultar, criar ou atualizar?", "Acesso e documentação", "Se aplicável"),
    ("Conhecimento", "Quais fontes são oficiais e qual hierarquia resolve conflitos?", "Fontes e proprietário", "Obrigatório"),
    ("Tom", "Quais exemplos representam a voz correta e quais devem ser evitados?", "Conversas reais", "Obrigatório"),
    ("Follow-up", "Quando retomar, quantas vezes e em quais horários?", "Cadência aprovada", "Se aplicável"),
    ("CRM", "Quais dados, eventos, tags e status devem ser registrados?", "Dicionário de dados", "Se aplicável"),
    ("Teste", "Quais cenários precisam passar antes do go-live?", "Bateria de testes", "Obrigatório"),
    ("Aceite", "Quem aprova e quais critérios determinam aprovação?", "Critérios assinados", "Obrigatório"),
    ("Governança", "Quem responde por conteúdo, operação, técnica e decisão?", "RACI", "Obrigatório"),
]
ws.append([])
ws.append(["Bloco", "Pergunta", "Evidência esperada", "Obrigatoriedade", "Resposta/Link", "Responsável", "Prazo", "Status"])
for item in onboarding:
    ws.append(item)
style_headers(ws, 4, 1, 8)
style_body(ws, 5, 22, 1, 8)
for row in range(5, 23):
    for col in range(5, 9):
        ws.cell(row, col).font = Font(name="Arial", size=10, color=BLUE_INPUT)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
dropdown(ws, "H5:H22", ["Não iniciado", "Aguardando cliente", "Em validação", "Aprovado", "Não aplicável"])
add_table(ws, "A4:H22", "tblOnboarding")
widths(ws, {"A": 18, "B": 65, "C": 32, "D": 18, "E": 42, "F": 22, "G": 15, "H": 22})
ws.freeze_panes = "A5"

ws = wb.create_sheet("Plano do Cliente")
title(ws, "Plano de Onboarding por agente", "Uma linha por agente. Esta aba transforma o briefing em especificação executável.")
headers = [
    "Cliente", "Agente", "Papel", "Canal", "Gatilho inicial", "Responsabilidades", "Fora de escopo",
    "Dados de entrada", "Saídas", "Regras de decisão", "Handoff humano", "Integrações", "Fonte oficial",
    "Dono Avraham", "Dono cliente", "Métrica principal", "Baseline", "Meta 90 dias", "Critério de aceite",
    "Go-live", "Status"
]
ws.append([])
ws.append(headers)
ws.append(["Cliente Exemplo", "Agente SDR", "Qualificar e encaminhar", "WhatsApp", "Resposta à prospecção", "Responder, qualificar, registrar e agendar", "Negociar desconto", "Nome, empresa, necessidade", "Lead qualificado + resumo", "Priorizar fit alto", "Transferir exceções", "CRM + Agenda", "Tabela comercial", "Líder IA", "Gerente comercial", "Leads qualificados", 120, 180, "10 cenários aprovados", "01/11/2026", "Planejado"])
style_headers(ws, 4, 1, len(headers))
style_body(ws, 5, 54, 1, len(headers))
for row in range(5, 55):
    for col in range(1, len(headers) + 1):
        ws.cell(row, col).font = Font(name="Arial", size=10, color=BLUE_INPUT)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
dropdown(ws, "U5:U54", ["Planejado", "Aguardando cliente", "Em construção", "Em teste", "Homologado", "Go-live", "Pausado"])
add_table(ws, "A4:U54", "tblPlanoCliente")
ws.freeze_panes = "D5"
for col in range(1, 22):
    ws.column_dimensions[ws.cell(4, col).column_letter].width = 20
for col in [6, 7, 8, 9, 10, 11, 12, 19]:
    ws.column_dimensions[ws.cell(4, col).column_letter].width = 30

ws = wb.create_sheet("Operação Mensal")
title(ws, "Custos e mensalidade", "Estime consumo e operação por cliente/mês. Preços de modelo e câmbio vêm da aba Parâmetros.")
headers = [
    "Cliente", "Mês", "Agentes ativos", "Conversas", "Turnos/conversa", "Tokens entrada/turno",
    "Tokens saída/turno", "Monitoramento (h)", "Otimização (h)", "Suporte (h)", "Relatório (h)",
    "Gestão cliente (h)", "Outras horas", "Total horas", "Custo trabalho", "Entrada (M tokens)",
    "Saída (M tokens)", "Custo tokens", "Canal/infra/terceiros", "Custo mensal total", "Mensalidade mínima",
    "Impacto conservador", "Referência por valor", "Mensalidade sugerida", "Mensalidade fechada", "Margem prevista"
]
ws.append([])
ws.append(headers)
ws.append(["Cliente Exemplo", "nov/2026", 1, 1200, 8, 700, 180, 2, 4, 4, 3, 2, 0, "=SUM(H5:M5)", "=N5*Parâmetros!$B$5", "=D5*E5*F5/1000000", "=D5*E5*G5/1000000", "=(P5*Parâmetros!$B$14+Q5*Parâmetros!$B$15)*Parâmetros!$B$13*(1+Parâmetros!$B$12)", 900, "=O5+R5+S5", "=T5/(1-Parâmetros!$B$7)", 25000, "=V5*Parâmetros!$B$11", "=MAX(U5,W5)", 6500, "=(Y5-T5)/Y5"])
for row in range(6, 55):
    ws.cell(row, 14, f'=IF(A{row}="","",SUM(H{row}:M{row}))')
    ws.cell(row, 15, f'=IF(A{row}="","",N{row}*Parâmetros!$B$5)')
    ws.cell(row, 16, f'=IF(A{row}="","",D{row}*E{row}*F{row}/1000000)')
    ws.cell(row, 17, f'=IF(A{row}="","",D{row}*E{row}*G{row}/1000000)')
    ws.cell(row, 18, f'=IF(A{row}="","",(P{row}*Parâmetros!$B$14+Q{row}*Parâmetros!$B$15)*Parâmetros!$B$13*(1+Parâmetros!$B$12))')
    ws.cell(row, 20, f'=IF(A{row}="","",O{row}+R{row}+S{row})')
    ws.cell(row, 21, f'=IF(A{row}="","",T{row}/(1-Parâmetros!$B$7))')
    ws.cell(row, 23, f'=IF(A{row}="","",V{row}*Parâmetros!$B$11)')
    ws.cell(row, 24, f'=IF(A{row}="","",MAX(U{row},W{row}))')
    ws.cell(row, 26, f'=IF(A{row}="","",(Y{row}-T{row})/Y{row})')
style_headers(ws, 4, 1, len(headers))
style_body(ws, 5, 54, 1, len(headers))
manual_cols = list(range(1, 14)) + [19, 22, 25]
for row in range(5, 55):
    for col in manual_cols:
        ws.cell(row, col).font = Font(name="Arial", size=10, color=BLUE_INPUT)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    for col in [15, 18, 19, 20, 21, 22, 23, 24, 25]:
        ws.cell(row, col).number_format = 'R$ #,##0.00;[Red](R$ #,##0.00);-'
    ws.cell(row, 26).number_format = "0.0%"
add_table(ws, "A4:Z54", "tblOperacaoMensal")
ws.freeze_panes = "D5"
for col in range(1, 27):
    ws.column_dimensions[ws.cell(4, col).column_letter].width = 18

ws = wb.create_sheet("Relatório ROI")
title(ws, "Relatório mensal de impacto e ROI", "Compare períodos equivalentes. Separe valores observados de estimativas e não conte o mesmo impacto duas vezes.")
headers = [
    "Cliente", "Mês", "Leads baseline", "Leads atual", "Resposta baseline (min)", "Resposta atual (min)",
    "Conversão baseline", "Conversão atual", "Vendas incrementais", "Ticket médio", "Margem contribuição",
    "Margem incremental", "Horas humanas baseline", "Horas humanas atual", "Horas liberadas", "Custo humano/h",
    "Economia trabalho", "Perdas evitadas", "Impacto total", "Setup pago", "Mensalidade", "Outros custos cliente",
    "Investimento do mês", "Benefício líquido", "ROI mensal", "Payback setup (meses)", "Tipo de evidência", "Fonte/observação"
]
ws.append([])
ws.append(headers)
ws.append(["Cliente Exemplo", "nov/2026", 1200, 1250, 180, 4, 0.08, 0.10, "=MAX(0,D5*H5-C5*G5)", 2500, 0.35, "=I5*J5*K5", 180, 120, "=MAX(0,M5-N5)", 55, "=O5*P5", 1000, "=L5+Q5+R5", 12000, 6500, 0, "=U5+V5", "=S5-W5", "=IF(W5=0,0,X5/W5)", "=IF(X5<=0,0,T5/X5)", "Observado", "CRM + folha de ponto"])
for row in range(6, 55):
    ws.cell(row, 9, f'=IF(A{row}="","",MAX(0,D{row}*H{row}-C{row}*G{row}))')
    ws.cell(row, 12, f'=IF(A{row}="","",I{row}*J{row}*K{row})')
    ws.cell(row, 15, f'=IF(A{row}="","",MAX(0,M{row}-N{row}))')
    ws.cell(row, 17, f'=IF(A{row}="","",O{row}*P{row})')
    ws.cell(row, 19, f'=IF(A{row}="","",L{row}+Q{row}+R{row})')
    ws.cell(row, 23, f'=IF(A{row}="","",U{row}+V{row})')
    ws.cell(row, 24, f'=IF(A{row}="","",S{row}-W{row})')
    ws.cell(row, 25, f'=IF(A{row}="","",IF(W{row}=0,0,X{row}/W{row}))')
    ws.cell(row, 26, f'=IF(A{row}="","",IF(X{row}<=0,0,T{row}/X{row}))')
style_headers(ws, 4, 1, len(headers))
style_body(ws, 5, 54, 1, len(headers))
manual_cols = list(range(1, 9)) + [10, 11, 13, 14, 16, 18, 20, 21, 22, 27, 28]
for row in range(5, 55):
    for col in manual_cols:
        ws.cell(row, col).font = Font(name="Arial", size=10, color=BLUE_INPUT)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    for col in [10, 12, 16, 17, 18, 19, 20, 21, 22, 23, 24]:
        ws.cell(row, col).number_format = 'R$ #,##0.00;[Red](R$ #,##0.00);-'
    for col in [7, 8, 11, 25]:
        ws.cell(row, col).number_format = "0.0%"
dropdown(ws, "AA5:AA54", ["Observado", "Estimado", "Contribuição provável", "Não mensurável"])
ws.conditional_formatting.add("Y5:Y54", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FECACA")))
ws.conditional_formatting.add("Y5:Y54", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=PALE_GREEN)))
add_table(ws, "A4:AB54", "tblRelatorioROI")
ws.freeze_panes = "D5"
for col in range(1, 29):
    ws.column_dimensions[ws.cell(4, col).column_letter].width = 18
ws.column_dimensions["AB"].width = 35

for sheet in wb.worksheets:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_view.zoomScale = 85
    for row in sheet.iter_rows():
        for cell in row:
            if cell.font.name != "Arial":
                cell.font = Font(name="Arial", size=cell.font.sz or 10, bold=cell.font.bold, italic=cell.font.italic, color=cell.font.color)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)

# Basic structural verification before external recalculation.
check = load_workbook(OUTPUT, data_only=False)
assert check.sheetnames == [
    "LEIA-ME", "Parâmetros", "Roteiro Pré-venda", "Oportunidades", "Escopo Setup",
    "Onboarding", "Plano do Cliente", "Operação Mensal", "Relatório ROI"
]
assert check["Oportunidades"]["Q5"].value.startswith("=")
assert check["Escopo Setup"]["AA5"].value.startswith("=")
assert check["Relatório ROI"]["Y5"].value.startswith("=")
print(OUTPUT)
